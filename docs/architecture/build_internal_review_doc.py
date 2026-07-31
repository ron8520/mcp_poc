from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "enterprise-mcp-platform-internal-review.md"
OUTPUT = HERE / "enterprise-mcp-platform-internal-review.docx"

FONT = "Calibri"
INK = "17324D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "526579"
LIGHT_GRAY = "E8EEF5"
PALE_BLUE = "EAF2F8"
GRID = "C7D0D9"
USABLE_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(run, *, name: str = FONT, size: float | None = None, color: str | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document) -> None:
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    caption = styles["Caption"]
    caption.font.name = FONT
    caption._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    caption.font.size = Pt(9)
    caption.font.color.rgb = rgb(MUTED)
    caption.font.italic = True
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def set_header_footer(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    left = p.add_run("ENTERPRISE MCP PLATFORM")
    set_run_font(left, size=8.5, color=MUTED, bold=True)
    right = p.add_run("\tDEVELOPER GUIDE")
    set_run_font(right, size=8.5, color=MUTED, bold=True)
    tabs = p.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(6.5))
    paragraph_border_bottom(p, GRID, "4")

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = fp.add_run("Developer Architecture Guide  |  ")
    set_run_font(label, size=8.5, color=MUTED)
    add_field(fp, "PAGE")
    tail = fp.add_run(" of ")
    set_run_font(tail, size=8.5, color=MUTED)
    add_field(fp, "NUMPAGES")


def paragraph_border_bottom(paragraph, color: str, size: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_geometry(table, widths: list[int], vertical_margin: int = 80) -> None:
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        set_cant_split(row)
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=vertical_margin, bottom=vertical_margin)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_pr.extend([r_fonts, color, underline])
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|https?://\S+)")


def add_inline(paragraph, text: str, *, base_size: float = 11, base_color: str = INK) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, size=base_size, color=base_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, color=base_color, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Courier New", size=max(8.5, base_size - 1), color=DARK_BLUE)
            shade_run(run, PALE_BLUE)
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        else:
            url = token.rstrip(".,;)")
            add_hyperlink(paragraph, url, url)
            if len(url) != len(token):
                run = paragraph.add_run(token[len(url) :])
                set_run_font(run, size=base_size, color=base_color)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=base_size, color=base_color)


def shade_run(run, fill: str) -> None:
    r_pr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    r_pr.append(shd)


def add_callout(doc: Document, label: str, text: str, *, fill: str = PALE_BLUE, accent: str = BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [USABLE_DXA])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=10.5, color=accent, bold=True)
    add_inline(p, text, base_size=10.5)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(96)
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("DEVELOPER ARCHITECTURE GUIDE")
    set_run_font(run, size=10.5, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    title_run = title.add_run("Enterprise MCP Platform on AWS")
    set_run_font(title_run, size=30, color=DARK_BLUE, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    subtitle_run = subtitle.add_run("Implementation, identity, policy and operations")
    set_run_font(subtitle_run, size=14, color=MUTED)

    for label, value in [
        ("Status", "Living technical guide"),
        ("Updated", "31 July 2026"),
        ("Maintainer", "Cloud Platform Team"),
        ("Implementation root", "examples/enterprise_mcp_platform"),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        lr = p.add_run(f"{label}: ")
        set_run_font(lr, size=10.5, color=INK, bold=True)
        vr = p.add_run(value)
        set_run_font(vr, name="Courier New" if label == "Implementation root" else FONT, size=10.5, color=INK)

    add_callout(
        doc,
        "Start here",
        "Use section 2 to find code, section 3 for local workflows, section 6 for the architecture, and section 17 for known gaps.",
    )
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(12)
    note.paragraph_format.space_after = Pt(0)
    nr = note.add_run("Editable diagrams are supplied in the accompanying draw.io files.")
    set_run_font(nr, size=9.5, color=MUTED, italic=True)
    doc.add_page_break()


def add_toc(doc: Document, source: str) -> None:
    doc.add_paragraph("Contents", style="Heading 1")
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(10)
    ir = intro.add_run("Developer reference sections")
    set_run_font(ir, size=9.5, color=MUTED, italic=True)
    for line in source.splitlines():
        if not line.startswith("## ") or line.startswith("## Developer Architecture Guide"):
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        add_inline(p, line[3:], base_size=10.2)
    doc.add_page_break()


def new_numbering_id(doc: Document, kind: str = "decimal") -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    number_format = "bullet" if kind == "bullet" else "decimal"
    level_text = "•" if kind == "bullet" else "%1."
    for tag, value in (("start", "1"), ("numFmt", number_format), ("lvlText", level_text), ("lvlJc", "left")):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:val"), value)
        lvl.append(node)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.extend([tabs, ind])
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "22")
    r_pr.extend([fonts, size])
    lvl.append(r_pr)
    abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(list(numbering).index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)


def table_widths(headers: list[str]) -> list[int]:
    count = len(headers)
    if count == 2:
        return [2800, 6560]
    if count == 3:
        if headers == ["Area", "State", "Evidence or gap"]:
            return [1900, 2300, 5160]
        if headers[0] in {"Tool", "Capability", "Decision"}:
            return [2500, 2900, 3960]
        return [2200, 3300, 3860]
    if count == 4:
        return [1550, 2750, 2600, 2460]
    if count == 5:
        return [1450, 1550, 1850, 1600, 2910]
    return [USABLE_DXA // count] * (count - 1) + [USABLE_DXA - (USABLE_DXA // count) * (count - 1)]


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    compact_controls = headers == ["Control", "Requirement", "Verification"]
    compact_reference = headers in [
        ["Area", "State", "Evidence or gap"],
        ["Component", "Responsibility", "Security boundary"],
    ]
    vertical_margin = 40 if compact_controls or compact_reference else 80
    body_size = 8.35 if compact_controls else 8.6 if compact_reference else 9.2
    header_size = 9.2 if compact_reference else 9.5
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, table_widths(headers), vertical_margin)
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(header) < 12 else WD_ALIGN_PARAGRAPH.LEFT
        add_inline(p, header, base_size=header_size)
        for run in p.runs:
            run.bold = True
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0 if compact_controls or compact_reference else 1.05
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(value) < 18 and idx != len(row_data) - 1 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, value, base_size=body_size)
    set_table_geometry(table, table_widths(headers), vertical_margin)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_threat_register(doc: Document, rows: list[list[str]]) -> None:
    add_callout(
        doc,
        "Reading the register",
        "Inherent and residual ratings are qualitative and provisional. Residual ratings assume every listed control is implemented and evidenced.",
    )
    bullet_id = new_numbering_id(doc, "bullet")
    for row in rows:
        threat_id, stride, threat, controls, inherent, residual, validation = row
        p = doc.add_paragraph(style="Heading 3")
        p.paragraph_format.keep_with_next = True
        add_inline(p, f"{threat_id} - {stride}: {threat}", base_size=11.5, base_color=DARK_BLUE)
        for label, value in [
            ("Required controls", controls),
            ("Risk", f"Inherent: {inherent}; residual: {residual}."),
            ("Validation / owner", validation),
        ]:
            bp = doc.add_paragraph()
            bp.paragraph_format.space_after = Pt(4)
            bp.paragraph_format.line_spacing = 1.25
            apply_numbering(bp, bullet_id)
            lr = bp.add_run(f"{label}: ")
            set_run_font(lr, size=10.2, color=INK, bold=True)
            add_inline(bp, value, base_size=10.2)


def add_code_block(doc: Document, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [USABLE_DXA])
    cell = table.cell(0, 0)
    shade_cell(cell, "F6F8FA")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run("\n".join(lines))
    set_run_font(run, name="Courier New", size=8.6, color=INK)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, alt: str, target: str) -> None:
    svg = (SOURCE.parent / target).resolve()
    png = svg.with_suffix(".png")
    if not png.exists():
        raise FileNotFoundError(f"rendered diagram missing: {png}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    figure_width = 5.75 if "target-architecture" in target else 5.9
    inline = run.add_picture(str(png), width=Inches(figure_width))
    doc_pr = inline._inline.docPr
    doc_pr.set("title", alt)
    doc_pr.set("descr", alt)


def parse_table(lines: list[str], start: int) -> tuple[list[str], list[list[str]], int]:
    raw = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        raw.append(lines[idx].strip())
        idx += 1
    parsed = [[cell.strip() for cell in row.strip("|").split("|")] for row in raw]
    headers = parsed[0]
    rows = parsed[2:]
    return headers, rows, idx


def parse_markdown(doc: Document, source: str) -> None:
    lines = source.splitlines()
    idx = 0
    in_code = False
    code_lines: list[str] = []
    paragraph_lines: list[str] = []
    skipped_source_header = False
    active_num_id: int | None = None
    active_bullet_id: int | None = None

    def flush_paragraph() -> None:
        nonlocal paragraph_lines, active_num_id
        if not paragraph_lines:
            return
        text = " ".join(line.strip() for line in paragraph_lines).strip()
        if text:
            p = doc.add_paragraph()
            add_inline(p, text)
        paragraph_lines = []
        active_num_id = None

    while idx < len(lines):
        line = lines[idx]
        stripped = line.strip()

        if active_num_id is not None and not re.match(r"^\d+\.\s+", stripped):
            active_num_id = None
        if active_bullet_id is not None and not stripped.startswith("- "):
            active_bullet_id = None

        if stripped.startswith("```"):
            flush_paragraph()
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            idx += 1
            continue
        if in_code:
            code_lines.append(line)
            idx += 1
            continue

        if not stripped:
            flush_paragraph()
            idx += 1
            continue

        if stripped == "---":
            flush_paragraph()
            idx += 1
            continue

        if stripped.startswith("# ") and not skipped_source_header:
            flush_paragraph()
            skipped_source_header = True
            idx += 1
            continue

        if stripped.startswith("## "):
            flush_paragraph()
            doc.add_paragraph(stripped[3:], style="Heading 1")
            idx += 1
            continue
        if stripped.startswith("### "):
            flush_paragraph()
            doc.add_paragraph(stripped[4:], style="Heading 2")
            idx += 1
            continue

        image_match = re.match(r"!\[([^\]]+)\]\(([^)]+)\)", stripped)
        if image_match:
            flush_paragraph()
            add_figure(doc, image_match.group(1), image_match.group(2))
            idx += 1
            continue

        if stripped.startswith("|") and idx + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[idx + 1].strip()):
            flush_paragraph()
            headers, rows, idx = parse_table(lines, idx)
            if headers[:2] == ["ID", "STRIDE"]:
                add_threat_register(doc, rows)
            else:
                add_table(doc, headers, rows)
            continue

        if stripped.startswith("- "):
            flush_paragraph()
            if active_bullet_id is None:
                active_bullet_id = new_numbering_id(doc, "bullet")
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            apply_numbering(p, active_bullet_id)
            add_inline(p, stripped[2:])
            idx += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.*)", stripped)
        if numbered:
            if paragraph_lines:
                flush_paragraph()
            if active_num_id is None:
                active_num_id = new_numbering_id(doc)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            apply_numbering(p, active_num_id)
            add_inline(p, numbered.group(1))
            idx += 1
            continue

        if stripped.startswith("**") and stripped.endswith("**") and stripped.count("**") == 2:
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            add_inline(p, stripped, base_size=11)
            idx += 1
            continue

        paragraph_lines.append(line)
        idx += 1

    flush_paragraph()


def set_document_settings(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    doc.core_properties.title = "Enterprise MCP Platform on AWS - Developer Architecture Guide"
    doc.core_properties.subject = "Developer reference for architecture, identity, policy, deployment and operations"
    doc.core_properties.author = "Cloud Platform Team"
    doc.core_properties.keywords = "MCP, AgentCore, Gateway, Runtime, Entra, SharePoint, developer guide, architecture"
    doc.core_properties.comments = "Generated from the repository developer-guide source."


def build() -> None:
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    set_header_footer(doc.sections[0])
    set_document_settings(doc)
    add_cover(doc)
    source = SOURCE.read_text(encoding="utf-8")
    add_toc(doc, source)
    body_start = source.index("## 1. Platform overview")
    parse_markdown(doc, source[body_start:])

    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
